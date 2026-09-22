#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Aplica no deck final (exportado do Canva) as correções da reconferência.

Regras que não se quebram aqui:

  * nunca atribuir ``text_frame.text`` — isso colapsa o parágrafo num run sem
    estilo e o slide perde fonte, cor e tamanho. Reescreve-se ``run.text``.
  * para criar uma forma nova, copiar o XML de uma existente e mudar só o que
    precisa: é o único jeito de herdar o design do Canva sem recriá-lo.
  * caixa do Canva vem com ``spAutoFit`` e largura feita para o texto que
    tinha. Texto mais largo quebra linha — desligar a quebra ou alargar.

Uso:  python3 ferramenta/corrigir_deck_final.py entrada.pptx saida.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
ENT = sys.argv[1] if len(sys.argv) > 1 else "entrada.pptx"
SAI = sys.argv[2] if len(sys.argv) > 2 else "saida.pptx"

FUNDO = "071A2F"          # o azul escuro do fundo dos cartões
mudancas = []


def log(texto):
    mudancas.append(texto)


def por_id(slide, ident):
    for sh in slide.shapes:
        if sh.shape_id == ident:
            return sh
    raise SystemExit(f"forma id={ident} não encontrada")


def texto(shape, novo, paragrafo=0):
    """Troca o texto de um parágrafo mantendo o estilo do primeiro run."""
    par = shape.text_frame.paragraphs[paragrafo]
    if not par.runs:
        raise SystemExit(f"parágrafo {paragrafo} sem run em id={shape.shape_id}")
    par.runs[0].text = novo
    for r in par.runs[1:]:
        r.text = ""


def geometria(shape, x=None, y=None, w=None, h=None, quebra=None):
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)
    if quebra is not None:
        shape.text_frame.word_wrap = quebra


def cor_do_texto(shape, rgb):
    """Pinta todos os runs da forma, inclusive os que o Canva fatiou."""
    for par in shape.text_frame.paragraphs:
        for r in par.runs:
            for fill in r._r.findall(f".//{NS}solidFill"):
                for clr in fill.findall(f"{NS}srgbClr"):
                    clr.set("val", rgb)


def so_um_paragrafo(shape):
    """Descarta todos os parágrafos menos o primeiro.

    Sem isto, clonar uma caixa que tinha cinco linhas devolve as cinco — foi o
    que aconteceu na primeira tentativa, e cada número do Nível AA apareceu
    arrastando a pilha inteira atrás de si.
    """
    corpo = shape.text_frame._txBody
    paragrafos = corpo.findall(f"{NS}p")
    for extra in paragrafos[1:]:
        corpo.remove(extra)


def clonar(slide, shape, x, y, novo_texto=None, w=None, quebra=None, tamanho=None):
    """Duplica uma forma preservando o XML e devolve a cópia já posicionada."""
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    copia = slide.shapes[-1]
    copia._element.nvSpPr.cNvPr.set("id", str(max(s.shape_id for s in slide.shapes) + 1))
    geometria(copia, x=x, y=y, w=w, quebra=quebra)
    if novo_texto is not None:
        so_um_paragrafo(copia)
        texto(copia, novo_texto)
    if tamanho is not None:
        for par in copia.text_frame.paragraphs:
            for r in par.runs:
                r.font.size = Pt(tamanho)
    return copia


def apagar(shape):
    shape._element.getparent().remove(shape._element)


pres = Presentation(ENT)
s1, s5, s6, s8, s9, s11, s12 = (pres.slides[i] for i in (0, 4, 5, 7, 8, 10, 11))

# ───────────────────────────────────────────── slide 1: falta um espaço
texto(por_id(s1, 10), "WCAG 2.1 · Portal do IFAL — Campus Maceió")
# a caixa fora dimensionada para o texto sem o espaço: um caractere a mais
# jogava "Maceió" para a segunda linha
geometria(por_id(s1, 10), quebra=False)
log("slide 1: 'WCAG2.1' → 'WCAG 2.1'")

# ───────────────────────── slides 5 e 6: a situação escrita na cor da pílula
# O Canva pintou a palavra ("Conforme", "Parcial"...) exatamente da mesma cor
# da pílula em que ela está: o texto existe, mas ninguém lê. Sobra só a cor —
# que é a falha 1.4.1 Uso de cores, a mesma que o deck acusa no banner do
# portal. Escurecer o texto resolve sem tirar a cor de ninguém.
palavras = ("Conforme", "Parcial", "Não conforme", "Não aplicável")
pintadas = 0
for slide, legenda_y in ((s5, 2.35), (s6, None)):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip().replace("  ", " ")
        # o Canva fatia a palavra em runs ("Parc" + " " + "ial")
        limpo = "".join(sh.text_frame.text.split())
        if limpo not in ("".join(p.split()) for p in palavras):
            continue
        if legenda_y is not None and abs(Emu(sh.top).inches - legenda_y) < 0.05:
            continue          # a legenda do topo fica como está: não tem pílula
        alvo = next(p for p in palavras if "".join(p.split()) == limpo)
        texto(sh, alvo)       # de quebra conserta o "Parc ial"
        cor_do_texto(sh, FUNDO)
        sh.text_frame.word_wrap = False
        pintadas += 1
log(f"slides 5 e 6: {pintadas} selos de situação agora legíveis sobre a pílula "
    f"(texto {FUNDO}); 'Parc ial' → 'Parcial'")

# slide 6: duas palavras coladas no item 9
texto(por_id(s6, 17), "Indicador de foco visível")
log("slide 6: 'focovisível' → 'foco visível'")

# ────────────────────── slide 8: sai a reprovação que era da marcação provisória
# O painel "CONTRASTE MEDIDO" trazia cinco linhas; a de 4,0:1 atribuída à Barra
# do Governo Federal media, na verdade, o cinza #7F7F7F do bloco de espera que
# o barra.js substitui. As quatro restantes redistribuem-se no mesmo espaço.
apagar(por_id(s8, 26))        # o "4,0:1"
apagar(por_id(s8, 23))        # o rótulo "Barra do Governo Federal"
LINHAS_CONTRASTE = [(2.60, 29, 24), (3.44, 28, 21), (4.28, 27, 20), (5.12, 18, 22)]
for y, id_valor, id_rotulo in LINHAS_CONTRASTE:
    geometria(por_id(s8, id_valor), y=y)
    geometria(por_id(s8, id_rotulo), y=y + 0.34)
log("slide 8: sai a linha '4,0:1 Barra do Governo Federal'; as outras quatro "
    "redistribuídas")

# ───────────────────────────────────────────────── slide 9: números e painel
texto(por_id(s9, 21), "3")
texto(por_id(s9, 24), "de 58 alvos < 24×24 px")
geometria(por_id(s9, 24), w=1.7, quebra=False)
log("slide 9: '4 reprovações' → 3; 'alvos < 24×24 px' → 'de 58 alvos < 24×24 px'")

# O painel da esquerda era todo sobre o reflow, que deixou de ser falha, e a
# imagem era a foto de outro slide, ilegível. Entra a tarefa complementar, que
# o enunciado pede e que não tinha slide nenhum.
rotulo, grande, linha = por_id(s9, 24), por_id(s9, 22), por_id(s9, 25)
for ident in (4, 16, 17, 18, 19):
    apagar(por_id(s9, ident))

clonar(s9, por_id(s9, 13), x=1.03, y=2.30, novo_texto="TAREFA COMPLEMENTAR · TALKBACK",
       w=3.6, quebra=False)
clonar(s9, grande, x=1.03, y=2.78, novo_texto="10 min 16 s", w=3.2, quebra=False)
clonar(s9, linha, x=1.03, y=3.42, w=4.4, quebra=True,
       novo_texto="para chegar ao primeiro edital aberto do campus, na 2ª tentativa")
clonar(s9, linha, x=1.03, y=4.10, w=4.4, quebra=True,
       novo_texto="1ª tentativa abandonada aos 12 min 4 s")
clonar(s9, linha, x=1.03, y=4.50, w=4.4, quebra=True,
       novo_texto="Samsung Galaxy A15 5G · Android 16 · TalkBack")
rodape = clonar(s9, linha, x=1.03, y=5.30, w=4.4, quebra=True,
                novo_texto="O que é do portal: a quantidade de elementos a percorrer "
                           "até o edital e a sequência de leitura pouco coerente com "
                           "a organização visual — critérios 1.3.2 e 2.4.3.")
for par in rodape.text_frame.paragraphs:           # nota de rodapé, peso menor
    for r in par.runs:
        r.font.size, r.font.bold = Pt(9), False
        cor_do_texto(rodape, "A9C1CF")

# O título falava do reflow, que saiu. Passa a cobrir as duas metades do slide:
# os números medidos, à direita, e o uso real, à esquerda.
texto(por_id(s9, 14), "No celular, o problema aparece de dois jeitos")
log("slide 9: painel do reflow (e a foto de slide) dá lugar à tarefa complementar; "
    "título ajustado ao novo conteúdo")

# ───────────────────────────────────────── slide 11: conformidade, as colunas
# Coluna NÍVEL A: entra o 1.3.1, que já estava como "Não conforme" no item 2 do
# checklist e não constava aqui. Três linhas, em ordem numérica.
LINHAS_A = [
    (3.35, "1.3.1", "Informações e relações",
     "Página de contato sem cabeçalho algum"),
    (4.53, "1.4.1", "Uso de cores",
     "Slide ativo diferenciado apenas pela cor"),
    (5.71, "2.2.2", "Pausar, Parar, Ocultar",
     "Banner avança a cada 4 s, sem pausa"),
]
num1, tit1, sub1 = por_id(s11, 7), por_id(s11, 18), por_id(s11, 22)
num2, tit2, sub2 = por_id(s11, 8), por_id(s11, 19), por_id(s11, 21)
trios = [(num1, tit1, sub1), (num2, tit2, sub2)]
trios.append((clonar(s11, num1, x=1.06, y=0), clonar(s11, tit1, x=1.98, y=0),
              clonar(s11, sub1, x=1.98, y=0)))
for (y, n, t, sub), (shn, sht, shs) in zip(LINHAS_A, trios):
    texto(shn, n)
    texto(sht, t)
    texto(shs, sub)
    geometria(shn, x=1.06, y=y, w=0.6, quebra=False)
    geometria(sht, x=1.98, y=y + 0.05, w=2.6, quebra=False)
    geometria(shs, x=1.98, y=y + 0.27, w=3.4, quebra=False)
log("slide 11, Nível A: entra 1.3.1; 'avançaacada 4s' → 'avança a cada 4 s'")

# Coluna NÍVEL AA: sai o 1.4.10 e as quatro linhas restantes se redistribuem.
# O número vinha de uma única caixa com as cinco linhas e espaços literais
# dentro ("1.4.1 0"). Vira uma caixa por linha: conserta o texto e desacopla a
# posição do espaçamento entre parágrafos.
LINHAS_AA = [
    (3.34, "1.4.3", 26, 31, "3 reprovações confirmadas; pior em 1,66:1"),
    (4.13, "1.4.11", 28, 33, "Foco em 1,56:1 sobre fundo branco"),
    (4.92, "1.4.12", 27, 32, "3 elementos cortados"),
    (5.70, "2.5.8", 25, 30, "23 de 58 alvos < 24×24 px"),
]
numeros = por_id(s11, 23)
for y, criterio, id_rot, id_val, valor in LINHAS_AA:
    rot, val = por_id(s11, id_rot), por_id(s11, id_val)
    delta = Emu(val.top).inches - Emu(rot.top).inches      # o desnível de cada par
    texto(val, valor)
    geometria(rot, y=y)
    geometria(val, y=y + delta, w=1.9, quebra=True)
    clonar(s11, numeros, x=7.18, y=y - 0.11, novo_texto=criterio, w=1.1, quebra=False)
apagar(numeros)
apagar(por_id(s11, 24))        # rótulo "Reflow"

# A caixa do "330 px em tela de 320 px" é reaproveitada como nota de rodapé do
# painel: o 2.5.8 é critério da WCAG 2.2 e o deck se intitula WCAG 2.1.
nota = por_id(s11, 29)
texto(nota, "O critério 2.5.8 pertence à WCAG 2.2 e entra como complemento — o "
            "enunciado admite \"2.1 ou mais recentes\". Na 2.1 o equivalente é o "
            "2.5.5, de Nível AAA, também não cumprido.")
geometria(nota, x=7.18, y=6.25, w=4.7, quebra=True)
log("slide 11, Nível AA: sai o 1.4.10; 4→3 reprovações; sai 'Rolagem horizontal +'; "
    "23 de 51→58; '1.4.1 0' e '1.4. 3' consertados; nota sobre a WCAG 2.2")

# ─────────────────────────────────────────────── slide 12: a recomendação que
# perdeu o objeto dá lugar ao 1.4.12, que é falha de AA no slide 11 e não tinha
# recomendação em lugar nenhum do deck.
texto(por_id(s12, 36), "Layout quebra com espaçamento de texto ampliado")
geometria(por_id(s12, 36), w=3.6, quebra=False)
texto(por_id(s12, 45), "1.4.12 (AA)")
geometria(por_id(s12, 45), quebra=False)   # mesma largura: o texto tem o mesmo tamanho
media = por_id(s12, 20)                       # o selo "MÉDIA" já existente
clonar(s12, media, x=Emu(media.left).inches, y=3.30, w=Emu(media.width).inches,
       quebra=False)
apagar(por_id(s12, 19))                       # o selo "ALTA" desta linha
log("slide 12: 'Rolagem horizontal em telas estreitas / 1.4.10 (AA)' → "
    "'Layout quebra com espaçamento de texto ampliado / 1.4.12 (AA)', prioridade MÉDIA")

pres.save(SAI)
print(f"salvo {SAI} — {len(pres.slides)} slides\n")
for m in mudancas:
    print("  ·", m)
